"""
extractor.py - Extrae texto plano de archivos PDF, DOCX y TXT.

Punto de entrada: extraer_texto(ruta) -> str
Cualquier problema de lectura se reporta como ExtraccionError, para que
main.py pueda anotarlo en el informe y continuar con el siguiente archivo.
"""
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from config import MAX_CARACTERES, PAGINAS_VISION


class ExtraccionError(Exception):
    """No se pudo obtener texto utilizable de un archivo."""


class SinTextoError(ExtraccionError):
    """PDF sin capa de texto (escaneado o hecho de imágenes): se puede intentar con visión."""


def extraer_texto(ruta: Path) -> str:
    """Devuelve el texto del archivo, normalizado y recortado a MAX_CARACTERES."""
    lectores = {".pdf": _leer_pdf, ".docx": _leer_docx, ".txt": _leer_txt}
    extension = ruta.suffix.lower()

    if extension not in lectores:
        raise ExtraccionError(f"Formato no soportado: {extension}")

    try:
        texto = lectores[extension](ruta)
    except ExtraccionError:
        raise
    except Exception as e:  # archivo corrupto, bloqueado, protegido, etc.
        raise ExtraccionError(f"No se pudo leer el archivo: {e}") from e

    # Colapsa saltos de línea y espacios múltiples: ahorra tokens sin perder significado.
    texto = re.sub(r"\s+", " ", texto).strip()

    if not texto:
        if extension == ".pdf":
            raise SinTextoError("El PDF no tiene texto seleccionable (es una imagen/escaneo).")
        raise ExtraccionError("El archivo no contiene texto utilizable.")

    return texto[:MAX_CARACTERES]


def _leer_pdf(ruta: Path) -> str:
    lector = PdfReader(ruta)

    if lector.is_encrypted:
        # Muchos PDF "protegidos" solo tienen contraseña de propietario (clave vacía).
        if not lector.decrypt(""):
            raise ExtraccionError("El PDF está protegido con contraseña.")

    partes, total = [], 0
    for pagina in lector.pages:
        texto_pagina = pagina.extract_text() or ""
        partes.append(texto_pagina)
        total += len(texto_pagina)
        if total >= MAX_CARACTERES:  # no hace falta leer un PDF de 300 páginas completo
            break
    return "\n".join(partes)


def _leer_docx(ruta: Path) -> str:
    documento = Document(str(ruta))
    partes = [p.text for p in documento.paragraphs if p.text.strip()]

    # Las tablas no aparecen en documento.paragraphs: se leen aparte.
    for tabla in documento.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(celda.text.strip() for celda in fila.cells))

    return "\n".join(partes)


def _leer_txt(ruta: Path) -> str:
    datos = ruta.read_bytes()
    # En Windows es común que un .txt esté en cp1252 y no en UTF-8.
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode("latin-1", errors="replace")  # último recurso: nunca falla


def renderizar_paginas_pdf(ruta: Path) -> list[bytes]:
    """Convierte las primeras páginas de un PDF en imágenes PNG (para el modelo de visión)."""
    try:
        import pymupdf  # se importa aquí: solo hace falta para PDF escaneados
    except ImportError as e:
        raise ExtraccionError(
            "Falta PyMuPDF para leer PDF escaneados. Ejecuta:  python -m pip install pymupdf"
        ) from e

    try:
        with pymupdf.open(str(ruta)) as documento:
            if documento.needs_pass:
                raise ExtraccionError("El PDF está protegido con contraseña.")
            zoom = pymupdf.Matrix(1.5, 1.5)  # ~110 dpi: legible y liviano
            return [
                documento[i].get_pixmap(matrix=zoom).tobytes("png")
                for i in range(min(PAGINAS_VISION, len(documento)))
            ]
    except ExtraccionError:
        raise
    except Exception as e:
        raise ExtraccionError(f"No se pudo convertir el PDF en imágenes: {e}") from e
